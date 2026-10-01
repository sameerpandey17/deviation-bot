"""
agent.py — LangGraph graph, tools, prompts, and file parsing helpers.

Graph flow:
  input (message, optional file_text, current_form)
    ↓
  [router]   → passes through (intent determined in extract)
    ↓
  [extract]  → single Groq call returns changed fields + risk as JSON
    ↓
  [validate] → checks dropdowns, date format, 2000-char limit
    ↓
  [merge]    → merges into current_form (edits never wipe other fields)
    ↓
  returns { form, risk, reply, missing_fields }
"""

import io
import json
import os
import re
from typing import Any, Dict, List, Optional, TypedDict

from dotenv import load_dotenv

load_dotenv()

# ══════════════════════════════════════════════════════════════════════════════
#  PROMPTS  (all in one place, top of file)
# ══════════════════════════════════════════════════════════════════════════════

SYSTEM_PROMPT = """You are AIVOA — an expert AI assistant for pharmaceutical deviation management.
You help operators log, review, and risk-assess manufacturing deviations.

You ALWAYS reply with a single, valid JSON object and nothing else.
The JSON must match this exact schema:

{
  "intent":  "log" | "edit" | "question",
  "form": {
    "site_plant":            string | null,
    "date_of_occurrence":    string (dd-mm-yyyy) | null,
    "title":                 string | null,
    "source":                "Manufacturing" | "QC Laboratory" | "Warehouse" | "Utilities" | "Audit" | "Supplier" | "Other" | null,
    "related_product":       string | null,
    "batch_lot_number":      string | null,
    "product_strength":      string (e.g. "500 mg", "250 mg/5 mL") | null,
    "product_amount":        string (e.g. "48 kg", "10,000 units") | null,
    "dosage_form":           "Tablet" | "Capsule" | "Injection" | "Syrup" | "Cream" | "Powder" | "Other" | null,
    "manufacturing_date":    string (mm-yyyy) | null,
    "expiry_date":           string (mm-yyyy) | null,
    "detailed_description":  string (max 2000 chars) | null,
    "initial_impact":        "Product Quality" | "Patient Safety" | "Regulatory Compliance" | "Data Integrity" | "EHS" | "No Impact" | null,
    "initial_severity":      "Minor" | "Major" | "Critical" | null
  },
  "risk": {
    "severity":         "Minor" | "Major" | "Critical" | null,
    "impact":           string | null,
    "reason":           string (2-3 sentences) | null,
    "next_action":      string | null,
    "containment":      string (one line) | null,
    "root_cause_hints": [string, string] (2-3 bullets) | null
  },
  "reply":          string,
  "missing_fields": [string]
}

Rules:
- intent="log": user is describing a new deviation event or uploading a report. Extract all fields, assess risk, and generate a concise summary in `reply` (e.g. "I've filled 7 of 9 fields from the document. I couldn't find the date of occurrence.").
- intent="edit": user is updating, correcting, or providing missing fields (e.g. "batch number is BMX24602 and size is 48 kg"). Return ONLY the changed/added fields in `form`, re-evaluate `risk`, and provide a clear confirmation in `reply` (e.g. "Updated the batch number to BMX24602.").
- intent="question": user is asking a question (e.g. "Why did you pick that severity?", "What is the containment plan?"). Do NOT change form or risk. Provide the direct answer in `reply`.
- Never guess missing values. Keep unknown fields as null.
- Dates must be dd-mm-yyyy format for date_of_occurrence; mm-yyyy for manufacturing_date and expiry_date.
- initial_severity and risk.severity must match.
- initial_impact and risk.impact must match.
"""

EXTRACT_USER_TEMPLATE = """Current form state:
{current_form}

Current risk assessment:
{current_risk}

Recent conversation history:
{history_section}

User message:
{message}

{file_section}

Analyze the request, decide intent ("log", "edit", or "question"), and return the complete JSON object."""

# ══════════════════════════════════════════════════════════════════════════════
#  VALID DROPDOWN VALUES
# ══════════════════════════════════════════════════════════════════════════════

VALID_SOURCES = {"Manufacturing", "QC Laboratory", "Warehouse", "Utilities", "Audit", "Supplier", "Other"}
VALID_IMPACTS = {"Product Quality", "Patient Safety", "Regulatory Compliance", "Data Integrity", "EHS", "No Impact"}
VALID_SEVERITIES = {"Minor", "Major", "Critical"}
DATE_RE = re.compile(r"^\d{2}-\d{2}-\d{4}$")

# ══════════════════════════════════════════════════════════════════════════════
#  FILE PARSING HELPERS
# ══════════════════════════════════════════════════════════════════════════════

def parse_pdf(file_bytes: bytes) -> str:
    """Extract text from PDF using PyMuPDF; fall back to OCR if sparse."""
    import fitz  # PyMuPDF
    doc = fitz.open(stream=file_bytes, filetype="pdf")
    pages_text = []
    for page in doc:
        pages_text.append(page.get_text())
    text = "\n".join(pages_text).strip()

    # If almost no text found, try OCR via pytesseract
    if len(text) < 50:
        text = _ocr_pdf(doc)
    doc.close()
    return text


def _ocr_pdf(doc) -> str:
    """Render each PDF page as image and OCR it."""
    try:
        import pytesseract
        from PIL import Image
        lines = []
        for page in doc:
            pix = page.get_pixmap(dpi=200)
            img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
            lines.append(pytesseract.image_to_string(img))
        return "\n".join(lines).strip()
    except Exception:
        return ""


def parse_docx(file_bytes: bytes) -> str:
    """Extract text from DOCX."""
    from docx import Document
    doc = Document(io.BytesIO(file_bytes))
    return "\n".join(p.text for p in doc.paragraphs).strip()


def parse_xlsx(file_bytes: bytes) -> str:
    """Extract text from XLSX (all sheets, all cells)."""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    rows = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                rows.append(" | ".join(cells))
    return "\n".join(rows)


def parse_image(file_bytes: bytes) -> str:
    """OCR an image file."""
    try:
        import pytesseract
        from PIL import Image
        img = Image.open(io.BytesIO(file_bytes))
        return pytesseract.image_to_string(img).strip()
    except Exception:
        return ""


def extract_text_from_file(filename: str, file_bytes: bytes) -> str:
    """Dispatch to the right parser based on file extension."""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext == "pdf":
        return parse_pdf(file_bytes)
    elif ext == "docx":
        return parse_docx(file_bytes)
    elif ext in ("xlsx", "xls"):
        return parse_xlsx(file_bytes)
    elif ext == "txt":
        return file_bytes.decode("utf-8", errors="replace")
    elif ext in ("jpg", "jpeg", "png"):
        return parse_image(file_bytes)
    else:
        return file_bytes.decode("utf-8", errors="replace")


# ══════════════════════════════════════════════════════════════════════════════
#  GROQ LLM CALL (direct, no LangChain messages dependency)
# ══════════════════════════════════════════════════════════════════════════════

def _call_groq(user_content: str) -> str:
    """Make a single Groq API call and return the raw text response.
    Tries primary requested model (llama-3.3-70b-versatile) first, then falls back
    to available high-capability models on this key (e.g. openai/gpt-oss-120b, openai/gpt-oss-20b).
    """
    from groq import Groq
    client = Groq(api_key=os.getenv("GROQ_API_KEY"))
    candidate_models = [
        "llama-3.3-70b-versatile",
        "openai/gpt-oss-120b",
        "openai/gpt-oss-20b",
        "qwen/qwen3.8-27b",
    ]
    last_err = None
    for model in candidate_models:
        try:
            response = client.chat.completions.create(
                model=model,
                messages=[
                    {"role": "system", "content": SYSTEM_PROMPT},
                    {"role": "user", "content": user_content},
                ],
                temperature=0,
                response_format={"type": "json_object"},
            )
            return response.choices[0].message.content.strip()
        except Exception as e:
            last_err = e
            continue
    raise RuntimeError(f"All Groq model attempts failed: {last_err}")


# ══════════════════════════════════════════════════════════════════════════════
#  LANGGRAPH STATE
# ══════════════════════════════════════════════════════════════════════════════

class AgentState(TypedDict):
    message: str
    file_text: Optional[str]
    current_form: Dict[str, Any]
    current_risk: Optional[Dict[str, Any]]
    history: Optional[List[Dict[str, str]]]
    # Populated by nodes:
    intent: Optional[str]
    raw_json: Optional[Dict]
    form: Optional[Dict]
    changed_fields: Optional[List[str]]
    risk: Optional[Dict]
    reply: Optional[str]
    missing_fields: Optional[List]


# ══════════════════════════════════════════════════════════════════════════════
#  GRAPH NODES
# ══════════════════════════════════════════════════════════════════════════════

def router_node(state: AgentState) -> AgentState:
    """Passes through — if a file is present, intent is log/extract."""
    if state.get("file_text"):
        state["intent"] = "log"
    return state


def extract_node(state: AgentState) -> AgentState:
    """Single Groq call: extract fields + assess risk + detect intent."""
    file_section = ""
    if state.get("file_text"):
        file_section = f"\nDocument content:\n---\n{state['file_text'][:6000]}\n---"

    # Format conversation history
    history_lines = []
    for h in (state.get("history") or [])[-6:]:
        role = "User" if h.get("role") == "user" else "Assistant"
        history_lines.append(f"{role}: {h.get('text', '')}")
    history_section = "\n".join(history_lines) if history_lines else "(No previous messages)"

    user_content = EXTRACT_USER_TEMPLATE.format(
        current_form=json.dumps(state.get("current_form") or {}, indent=2),
        current_risk=json.dumps(state.get("current_risk") or {}, indent=2),
        history_section=history_section,
        message=state["message"] or "(File uploaded without text message)",
        file_section=file_section,
    )

    raw_text = _call_groq(user_content)

    # Strip markdown fences if present
    if raw_text.startswith("```"):
        raw_text = re.sub(r"^```[a-z]*\n?", "", raw_text)
        raw_text = re.sub(r"\n?```$", "", raw_text)

    try:
        parsed = json.loads(raw_text)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", raw_text, re.DOTALL)
        if match:
            try:
                parsed = json.loads(match.group())
            except json.JSONDecodeError:
                parsed = _fallback_response(raw_text)
        else:
            parsed = _fallback_response(raw_text)

    return {**state, "raw_json": parsed, "intent": parsed.get("intent", "question")}


def _fallback_response(raw_text: str) -> Dict:
    return {
        "intent": "question",
        "form": None,
        "risk": None,
        "reply": raw_text or "I encountered an issue processing your request. Please try again.",
        "missing_fields": [],
    }


def validate_node(state: AgentState) -> AgentState:
    """Validate dropdown values, date format, char limits. Invalid → null."""
    parsed = dict(state.get("raw_json") or {})
    form = dict(parsed.get("form") or {})
    risk = dict(parsed.get("risk") or {})

    # Validate form dropdowns
    if form.get("source") and form["source"] not in VALID_SOURCES:
        form["source"] = None
    if form.get("initial_impact") and form["initial_impact"] not in VALID_IMPACTS:
        form["initial_impact"] = None
    if form.get("initial_severity") and form["initial_severity"] not in VALID_SEVERITIES:
        form["initial_severity"] = None

    # Validate date format
    if form.get("date_of_occurrence") and not DATE_RE.match(str(form["date_of_occurrence"])):
        form["date_of_occurrence"] = None

    # Validate mm-yyyy dates
    MONTH_YEAR_RE = re.compile(r"^\d{2}-\d{4}$")
    for date_field in ("manufacturing_date", "expiry_date"):
        val = form.get(date_field)
        if val and not MONTH_YEAR_RE.match(str(val)):
            form[date_field] = None

    # Enforce 2000-char limit
    if form.get("detailed_description") and len(form["detailed_description"]) > 2000:
        form["detailed_description"] = form["detailed_description"][:2000]

    # Validate risk dropdowns
    if risk.get("severity") and risk["severity"] not in VALID_SEVERITIES:
        risk["severity"] = None

    # Keep severity in sync between form and risk
    if form.get("initial_severity") and not risk.get("severity"):
        risk["severity"] = form["initial_severity"]
    elif risk.get("severity") and not form.get("initial_severity"):
        form["initial_severity"] = risk["severity"]

    # Keep impact in sync
    if form.get("initial_impact") and not risk.get("impact"):
        risk["impact"] = form["initial_impact"]
    elif risk.get("impact") and not form.get("initial_impact"):
        form["initial_impact"] = risk["impact"]

    return {**state, "raw_json": {**parsed, "form": form, "risk": risk}}


def merge_node(state: AgentState) -> AgentState:
    """Merge AI-returned fields into current_form; never wipe existing fields. Compute changed_fields."""
    parsed = state.get("raw_json") or {}
    intent = state.get("intent", "question")
    current_form = dict(state.get("current_form") or {})
    current_risk = dict(state.get("current_risk") or {})
    ai_form = dict(parsed.get("form") or {})
    ai_risk = parsed.get("risk")

    changed_fields = []
    if intent in ("log", "edit"):
        for key, value in ai_form.items():
            if value is not None and value != "":
                if current_form.get(key) != value:
                    changed_fields.append(key)
                current_form[key] = value
        final_risk = ai_risk if ai_risk else current_risk
    else:
        # question: keep form and risk unchanged
        final_risk = current_risk

    # Compute missing required fields
    required_keys = ["title", "date_of_occurrence", "source", "detailed_description", "initial_impact", "initial_severity"]
    missing = [k for k in required_keys if not current_form.get(k)]

    return {
        **state,
        "intent": intent,
        "form": current_form,
        "changed_fields": changed_fields,
        "risk": final_risk,
        "reply": parsed.get("reply", ""),
        "missing_fields": missing,
    }


# ══════════════════════════════════════════════════════════════════════════════
#  BUILD LANGGRAPH  (LangGraph 1.x API)
# ══════════════════════════════════════════════════════════════════════════════

def build_graph():
    from langgraph.graph import StateGraph, END

    graph = StateGraph(AgentState)
    graph.add_node("router", router_node)
    graph.add_node("extract", extract_node)
    graph.add_node("validate", validate_node)
    graph.add_node("merge", merge_node)

    graph.set_entry_point("router")
    graph.add_edge("router", "extract")
    graph.add_edge("extract", "validate")
    graph.add_edge("validate", "merge")
    graph.add_edge("merge", END)

    return graph.compile()


_graph = None

def get_graph():
    global _graph
    if _graph is None:
        _graph = build_graph()
    return _graph


# ══════════════════════════════════════════════════════════════════════════════
#  PUBLIC API
# ══════════════════════════════════════════════════════════════════════════════

def run_agent(
    message: str,
    current_form: Optional[Dict] = None,
    current_risk: Optional[Dict] = None,
    history: Optional[List[Dict]] = None,
    file_text: Optional[str] = None,
) -> Dict:
    """
    Run the LangGraph pipeline and return:
      { intent, form, changed_fields, risk, reply, missing_fields }
    """
    graph = get_graph()
    initial_state: AgentState = {
        "message": message,
        "file_text": file_text,
        "current_form": current_form or {},
        "current_risk": current_risk or {},
        "history": history or [],
        "intent": None,
        "raw_json": None,
        "form": None,
        "changed_fields": None,
        "risk": None,
        "reply": None,
        "missing_fields": None,
    }
    result = graph.invoke(initial_state)
    return {
        "intent": result.get("intent") or "log",
        "form": result.get("form") or {},
        "changed_fields": result.get("changed_fields") or [],
        "risk": result.get("risk") or {},
        "reply": result.get("reply") or "",
        "missing_fields": result.get("missing_fields") or [],
    }
